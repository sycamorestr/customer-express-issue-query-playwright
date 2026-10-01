"""Offline regressions for the unchanged workbook checkpoint/export processor.

Run with Python 3.10+ and openpyxl, or the package's portable Python:
    python -m unittest discover -s tests -p test_excel.py -v
All workbooks, results and attachments are synthetic temporary fixtures.
"""
import argparse
import copy
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.styles import Font, PatternFill


SCRIPT = (Path(__file__).resolve().parents[1] / "skills" /
          "customer-express-issue-query-playwright" / "scripts" / "excel.py")
SPEC = importlib.util.spec_from_file_location("workbook_processor", SCRIPT)
excel = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(excel)
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def replace_zip_parts(path, changes):
    """Rewrite fixture bytes without relying on production XML helpers."""
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        comment = archive.comment
    with zipfile.ZipFile(path, "w") as archive:
        archive.comment = comment
        for info, contents in entries:
            archive.writestr(info, changes.get(info.filename, contents))


class ExcelRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="express-excel-test-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "synthetic.xlsx"
        self.work = self.root / "checkpoint"
        self.output = self.root / "result.xlsx"

    def create_workbook(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "主表"
        ws["A3"], ws["B3"], ws["C3"] = "说明", "月结赔付单号", "快递单号"
        ws["B4"] = "圆通: @yt12345678；极兔 jt87654321"
        ws["B5"] = "@yt12345678, YT12345678"
        ws["B6"] = "ZT12345678"
        ws["C4"] = "this alternate column must not be queried"
        ws["F4"] = "=SUM(1,2)"
        ws["F4"].number_format = "0.00"
        ws["F4"].font = Font(bold=True, color="FF123456")
        ws["M8"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
        ws.merge_cells("I1:P1")
        ws["I1"] = "Original merged title"
        ws.column_dimensions["D"].width = 28
        ws.row_dimensions[4].height = 36
        ws.freeze_panes = "C4"
        other = wb.create_sheet("补充表")
        other.append(["物流单号", "原备注"])
        other.append(["YT12345678", "跨表重复"])
        other.append(["SF12345678", "独立单号"])
        info = wb.create_sheet("说明页")
        info["A1"] = "No tracking header"
        info["B1"] = "=1+1"
        wb.save(self.source)
        wb.close()
        # Simulate a producer that understates dimensions. Preparation must
        # inspect the actual rows, including rows after the declared boundary.
        with zipfile.ZipFile(self.source) as archive:
            sheet = archive.read("xl/worksheets/sheet1.xml")
        replace_zip_parts(self.source, {
            "xl/worksheets/sheet1.xml": sheet.replace(b'ref="A1:P8"', b'ref="A1:A1"', 1)
        })
        with zipfile.ZipFile(self.source, "a") as archive:
            archive.comment = b"preserve this archive comment"
            archive.writestr("xl/embeddings/synthetic-attachment.bin", b"\x00synthetic attachment\xff")
            archive.writestr("customXml/item1.xml", b"<metadata>synthetic only</metadata>")

    def prepare(self, *, column=None, sheets=None):
        return excel.prepare(argparse.Namespace(
            input=str(self.source), workdir=str(self.work), column=column, sheets=sheets))

    def valid_results(self):
        return {
            "YT12345678": {"no": "YT12345678", "matched": 2, "data": [
                ["YT12345678", "2026-01-01 12:00:00", "0", "=SUM(1,1)"],
                ["YT12345678", "2026-01-02 12:00:00", "19.90", "second order"],
            ]},
            "JT87654321": {"no": "JT87654321", "matched": 1, "data": [
                ["JT87654321", "2026-01-03 12:00:00", "5", "中文备注 & <text>"]
            ]},
            "ZT12345678": {"no": "ZT12345678", "matched": 0, "data": [
                ["ZT12345678", "", "", ""]
            ]},
            "SF12345678": {"no": "SF12345678", "matched": 1, "data": [
                ["SF12345678", "2026-01-04 12:00:00", "8", ""]
            ]},
        }

    def export(self):
        return excel.export(argparse.Namespace(workdir=str(self.work), output=str(self.output)))

    def ready_checkpoint(self):
        self.create_workbook()
        self.prepare()
        save_json(self.work / "results.json", self.valid_results())

    def assert_no_export(self):
        self.assertFalse(self.output.exists())
        self.assertFalse((self.work / "summary.json").exists())
        self.assertEqual(list(self.root.glob("excel-export-*.xlsx")), [])

    def test_round_trip_preserves_workbook_and_all_tracking_occurrences(self):
        self.ready_checkpoint()
        source_bytes = self.source.read_bytes()
        manifest = read_json(self.work / "manifest.json")
        self.assertEqual(manifest["unique"], ["YT12345678", "JT87654321", "ZT12345678", "SF12345678"])
        self.assertEqual(manifest["sheets"][0]["rows"], [
            {"row": 4, "nos": ["YT12345678", "JT87654321"]},
            {"row": 5, "nos": ["YT12345678", "YT12345678"]},
            {"row": 6, "nos": ["ZT12345678"]},
        ])
        self.assertEqual(manifest["skippedSheets"][0]["name"], "说明页")
        with zipfile.ZipFile(self.source) as original:
            self.assertEqual(ET.fromstring(original.read("xl/worksheets/sheet1.xml")).find("s:dimension", NS).get("ref"), "A1:A1")

        summary = self.export()
        self.assertEqual(self.source.read_bytes(), source_bytes)
        self.assertEqual((summary["unique"], summary["matched"], summary["notFound"], summary["multipleOrders"]), (4, 3, 1, 1))
        self.assertEqual(summary["sheets"][0]["resultColumns"], "Q:U")
        self.assertEqual(summary["sheets"][1]["resultColumns"], "C:G")
        self.assertEqual(summary["sheets"][0]["rows"], 3)
        self.assertEqual(summary["sheets"][1]["rows"], 2)
        with zipfile.ZipFile(self.source) as original, zipfile.ZipFile(self.output) as exported:
            self.assertEqual(original.namelist(), exported.namelist())
            self.assertEqual(original.comment, exported.comment)
            changed_sheets = {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"}
            for name in original.namelist():
                if name not in changed_sheets:
                    self.assertEqual(original.read(name), exported.read(name), name)
                else:
                    old_cells = ET.fromstring(original.read(name)).findall(".//s:c", NS)
                    new_cells = {cell.get("r"): cell for cell in ET.fromstring(exported.read(name)).findall(".//s:c", NS)}
                    for cell in old_cells:
                        self.assertEqual(ET.tostring(cell), ET.tostring(new_cells[cell.get("r")]), cell.get("r"))
        wb = openpyxl.load_workbook(self.output, data_only=False)
        try:
            ws = wb["主表"]
            self.assertEqual(ws["B4"].value, "圆通: @yt12345678；极兔 jt87654321")
            self.assertEqual(ws["F4"].value, "=SUM(1,2)")
            self.assertEqual(ws["F4"].data_type, "f")
            self.assertEqual(ws["F4"].number_format, "0.00")
            self.assertTrue(ws["F4"].font.bold)
            self.assertEqual(ws["M8"].fill.fgColor.rgb, "FFFFFF00")
            self.assertIsNone(ws["M8"].value)
            self.assertEqual(str(ws.merged_cells), "I1:P1")
            self.assertEqual(ws.column_dimensions["D"].width, 28)
            self.assertEqual(ws.row_dimensions[4].height, 36)
            self.assertEqual(ws.freeze_panes, "C4")
            self.assertEqual(ws["Q4"].value, "YT12345678\nYT12345678\nJT87654321")
            self.assertEqual(ws["S4"].value, "0\n19.90\n5")
            self.assertEqual(ws["T4"].value, "=SUM(1,1)\nsecond order\n中文备注 & <text>")
            self.assertEqual(ws["T4"].data_type, "s")
            self.assertEqual(ws["S5"].value, "0\n19.90\n0\n19.90")
            self.assertEqual(ws["U6"].value, "未找到订单")
            self.assertEqual(wb["补充表"]["C2"].value, "YT12345678\nYT12345678")
        finally:
            wb.close()

    def test_partial_malformed_and_unsafe_results_cannot_be_exported(self):
        self.ready_checkpoint()
        valid = self.valid_results()
        variants = {}
        partial = copy.deepcopy(valid)
        del partial["SF12345678"]
        variants["incomplete"] = partial
        mismatched_count = copy.deepcopy(valid)
        mismatched_count["YT12345678"]["matched"] = 1
        variants["wrong count"] = mismatched_count
        numeric_amount = copy.deepcopy(valid)
        numeric_amount["YT12345678"]["data"][0][2] = 0
        variants["non-string data"] = numeric_amount
        false_match = copy.deepcopy(valid)
        false_match["ZT12345678"]["data"][0][3] = "should be blank"
        variants["unmatched with data"] = false_match
        bad_xml = copy.deepcopy(valid)
        bad_xml["YT12345678"]["data"][0][3] = "invalid\x01character"
        variants["XML control character"] = bad_xml
        too_long = copy.deepcopy(valid)
        too_long["YT12345678"]["data"][0][3] = "x" * 32768
        variants["Excel text limit"] = too_long
        for label, results in variants.items():
            with self.subTest(label=label):
                save_json(self.work / "results.json", results)
                with self.assertRaises(ValueError):
                    self.export()
                self.assert_no_export()

    def test_modified_source_or_manifest_cannot_be_exported(self):
        self.ready_checkpoint()
        manifest_path = self.work / "manifest.json"
        manifest = read_json(manifest_path)
        modified = copy.deepcopy(manifest)
        modified["sheets"][0]["rows"][0]["nos"] = ["SF12345678"]
        save_json(manifest_path, modified)
        with self.assertRaisesRegex(ValueError, "Manifest row mapping"):
            self.export()
        self.assert_no_export()
        save_json(manifest_path, manifest)
        with zipfile.ZipFile(self.source, "a") as archive:
            archive.writestr("customXml/changed.xml", "source changed since preparation")
        with self.assertRaisesRegex(ValueError, "Source hash differs"):
            self.export()
        self.assert_no_export()

    def test_existing_checkpoint_and_output_are_not_overwritten(self):
        self.ready_checkpoint()
        manifest_bytes = (self.work / "manifest.json").read_bytes()
        with self.assertRaisesRegex(ValueError, "checkpoint already exists"):
            self.prepare()
        self.assertEqual((self.work / "manifest.json").read_bytes(), manifest_bytes)
        self.output.write_bytes(b"existing output must survive")
        with self.assertRaisesRegex(ValueError, "Refusing to overwrite output"):
            self.export()
        self.assertEqual(self.output.read_bytes(), b"existing output must survive")
        self.assertFalse((self.work / "summary.json").exists())

    def test_invalid_tracking_cell_prevents_checkpoint_creation(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["快递单号"])
        ws.append(["YT12345678"])
        ws.append(["=12345678"])
        ws.append([1234567890123456])
        ws.append([True])
        wb.save(self.source)
        wb.close()
        with self.assertRaisesRegex(ValueError, "Preparation failed"):
            self.prepare()
        self.assertFalse((self.work / "manifest.json").exists())
        report = read_json(self.work / "prepare-errors.json")
        self.assertEqual([error["row"] for error in report["errors"]], [3, 4, 5])

    def test_explicit_selection_requires_a_header_in_every_selected_sheet(self):
        self.create_workbook()
        with self.assertRaisesRegex(ValueError, "Preparation failed"):
            self.prepare(sheets="主表,说明页")
        self.assertFalse((self.work / "manifest.json").exists())
        report = read_json(self.work / "prepare-errors.json")
        self.assertEqual(report["errors"][0]["sheet"], "说明页")


if __name__ == "__main__":
    unittest.main()
