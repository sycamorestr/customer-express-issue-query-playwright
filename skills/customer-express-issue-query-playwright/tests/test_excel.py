"""Offline regressions for safe preparation and lossless XLSX export.

Run with the skill's Python: python -B -m unittest discover -s tests -v
All workbooks and identifiers in this file are synthetic.
"""
import argparse
import hashlib
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


SKILL_ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("skill_excel", SKILL_ROOT / "scripts" / "excel.py")
excel = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(excel)
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ExcelRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "synthetic.xlsx"
        self.work = self.root / "work"

    def make_workbook(self, rows=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Courier A"
        for row in rows or [["单号", "原备注"], ["AB12345678", "原始说明"]]:
            ws.append(row)
        wb.save(self.source)
        wb.close()

    def prepare(self, column=None, sheets=None):
        return excel.prepare(argparse.Namespace(input=str(self.source), workdir=str(self.work),
                                                column=column, sheets=sheets))

    def manifest(self):
        return json.loads((self.work / "manifest.json").read_text(encoding="utf-8"))

    def write_results(self, values=None):
        if values is None:
            values = {no: {"no": no, "matched": 1,
                           "data": [[no, "2026-09-01 12:00:00", "0", "Synthetic remark"]]}
                      for no in self.manifest()["unique"]}
        (self.work / "results.json").write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")

    def export(self, output=None):
        return excel.export(argparse.Namespace(workdir=str(self.work),
                                               output=str(output) if output else None))

    def package(self, path):
        with zipfile.ZipFile(path) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def replace_parts(self, replacements):
        parts = self.package(self.source)
        parts.update(replacements)
        with zipfile.ZipFile(self.source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.comment = b"synthetic preservation fixture"
            for name, content in parts.items():
                archive.writestr(name, content)

    def effective_column(self, sheet_xml, index):
        """Resolve column attributes as an Excel reader does for a column index."""
        matches = [dict(col.attrib) for col in ET.fromstring(sheet_xml).findall("s:cols/s:col", NS)
                   if int(col.get("min")) <= index <= int(col.get("max"))]
        if not matches:
            return {}
        result = matches[-1]
        return {key: value for key, value in result.items() if key not in ("min", "max")}

    def test_all_nine_sheets_include_generic_header_fallback(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for index in range(9):
            ws = wb.create_sheet(f"Courier {index + 1}")
            header = "月结赔付单号" if index == 0 else "单号"
            ws.append(["说明", "", ""])
            ws.append(["原记录", "描述", header])
            ws.append([index, "Synthetic", f"ZX{index + 10000000}"])
        wb.create_sheet("说明").append(["Summary information only"])
        wb.save(self.source)
        wb.close()
        self.prepare()
        manifest = self.manifest()
        self.assertEqual(len(manifest["sheets"]), 9)
        self.assertEqual(len(manifest["unique"]), 9)
        self.assertEqual([sheet["headerRow"] for sheet in manifest["sheets"]], [2] * 9)
        self.assertEqual([sheet["column"] for sheet in manifest["sheets"]], [3] * 9)
        self.assertEqual([sheet["name"] for sheet in manifest["skippedSheets"]], ["说明"])

    def test_explicit_headers_take_priority_over_generic_header(self):
        for preferred in ("月结赔付单号", "快递单号", "物流单号", "运单号", "发货单号"):
            with self.subTest(preferred=preferred), tempfile.TemporaryDirectory(dir=self.root) as directory:
                self.work = Path(directory) / "work"
                self.make_workbook([["单号", preferred], ["UNKNOWN", "ZX12345678"]])
                self.prepare()
                self.assertEqual(self.manifest()["unique"], ["ZX12345678"])
                self.assertEqual(self.manifest()["sheets"][0]["column"], 2)

    def test_duplicate_generic_headers_fail_with_auditable_candidates(self):
        self.make_workbook([["单号", "单号"], ["AB12345678", "CD12345678"]])
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertFalse((self.work / "manifest.json").exists())
        errors = json.loads((self.work / "prepare-errors.json").read_text(encoding="utf-8"))["errors"]
        self.assertTrue(any(len(error.get("candidates", [])) == 2 for error in errors))

    def test_header_must_be_exact_and_explicit_column_can_resolve_ambiguity(self):
        self.make_workbook([["单号备注", "订单号"], ["AB12345678", "CD12345678"]])
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertFalse((self.work / "manifest.json").exists())
        self.work = self.root / "ambiguous"
        self.make_workbook([["单号备注", "订单号", "快递单号", "物流单号"],
                            ["AB12345678", "CD12345678", "EF12345678", "GH12345678"]])
        with self.assertRaises(ValueError):
            self.prepare()
        self.work = self.root / "explicit"
        self.prepare(column="快递单号")
        self.assertEqual(self.manifest()["unique"], ["EF12345678"])
        self.assertEqual(self.manifest()["sheets"][0]["column"], 3)

    def test_boundary_separators_normalize_keys_without_changing_source(self):
        raw = " ，;|、 @ab12345678,， 极兔：cd12345678；\nEF12345678 |、;， "
        self.make_workbook([["单号"], [raw], ["ab12345678"]])
        original = sha256(self.source)
        self.prepare()
        manifest = self.manifest()
        self.assertEqual(manifest["unique"], ["AB12345678", "CD12345678", "EF12345678"])
        self.assertEqual(manifest["sheets"][0]["rows"][0]["nos"], manifest["unique"])
        self.assertEqual(manifest["sheets"][0]["rows"][1]["nos"], ["AB12345678"])
        self.assertEqual(sha256(self.source), original)
        wb = openpyxl.load_workbook(self.source)
        self.assertEqual(wb.active["A2"].value, raw)
        wb.close()

    def test_unsafe_tracking_values_produce_errors_instead_of_partial_manifest(self):
        values = ["AB12345678（待确认）", "备注:AB12345678", "AB1234", "AB12345678/",
                  "=12345678+1", "#VALUE!", 1234567890123456, 12345678.5,
                  True, ";，|", "@@AB12345678", "AB12 3456", "(AB12345678)"]
        self.make_workbook([["单号"], ["VALID12345678"], *[[value] for value in values]])
        with self.assertRaises(ValueError):
            self.prepare()
        errors = json.loads((self.work / "prepare-errors.json").read_text(encoding="utf-8"))["errors"]
        self.assertEqual({error["row"] for error in errors if "row" in error}, set(range(3, 3 + len(values))))
        self.assertFalse((self.work / "manifest.json").exists())

    def test_lossless_export_preserves_original_cells_structure_and_zip_parts(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Courier A"
        ws.append(["单号", "原金额", "原公式", "原备注"])
        ws.append(["AB12345678", 12.34, "=B2*2", "Original note"])
        ws["A2"].font = Font(bold=True, color="FF112233")
        ws["B2"].number_format = "0.00"
        ws["C2"].fill = PatternFill("solid", fgColor="FFABCDEF")
        ws["D2"].alignment = Alignment(wrap_text=True)
        ws["D2"].comment = Comment("Synthetic comment", "Fixture")
        ws["D2"].hyperlink = "https://example.invalid/synthetic"
        ws.merge_cells("E1:F1")
        ws["E1"] = "Merged original"
        ws["H4"].fill = PatternFill("solid", fgColor="FFCCFFCC")
        ws.column_dimensions["A"].width = 29
        ws.column_dimensions["B"].width = 13.5
        ws.row_dimensions[2].height = 31.5
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = "A1:D2"
        wb.create_sheet("Notes")["A1"] = "Untouched sheet"
        wb.save(self.source)
        wb.close()
        self.replace_parts({"xl/embeddings/synthetic-attachment.bin": b"attachment\x00\x01\xff",
                            "xl/cellimages.xml": b'<synthetic xmlns="urn:fixture">opaque image metadata</synthetic>'})
        original_hash = sha256(self.source)
        before = self.package(self.source)
        self.prepare()
        self.write_results()
        summary = self.export()
        after = self.package(summary["output"])
        changed_sheet = "xl/worksheets/sheet1.xml"
        self.assertEqual(before.keys(), after.keys())
        for name in before:
            if name != changed_sheet:
                self.assertEqual(before[name], after[name], name)
        old_root, new_root = ET.fromstring(before[changed_sheet]), ET.fromstring(after[changed_sheet])
        new_cells = {cell.get("r"): cell for cell in new_root.findall("s:sheetData/s:row/s:c", NS)}
        for cell in old_root.findall("s:sheetData/s:row/s:c", NS):
            self.assertEqual(ET.tostring(cell), ET.tostring(new_cells[cell.get("r")]), cell.get("r"))
        for tag in ("mergeCells", "sheetViews", "autoFilter", "sheetFormatPr", "hyperlinks", "pageMargins"):
            self.assertEqual(ET.tostring(old_root.find("s:" + tag, NS)),
                             ET.tostring(new_root.find("s:" + tag, NS)), tag)
        for row in old_root.findall("s:sheetData/s:row", NS):
            actual = new_root.find(f's:sheetData/s:row[@r="{row.get("r")}"]', NS)
            self.assertEqual(row.attrib, actual.attrib)
        for index in range(1, 9):
            self.assertEqual(self.effective_column(before[changed_sheet], index),
                             self.effective_column(after[changed_sheet], index))
        self.assertEqual(summary["sheets"][0]["resultColumns"], "I:M")
        self.assertEqual(sha256(self.source), original_hash)
        with zipfile.ZipFile(self.source) as old, zipfile.ZipFile(summary["output"]) as new:
            self.assertEqual(old.comment, new.comment)

    def test_result_columns_are_readable_without_changing_existing_dimensions(self):
        self.make_workbook()
        self.prepare()
        self.write_results()
        summary = self.export()
        wb = openpyxl.load_workbook(summary["output"])
        ws = wb.active
        for letter in "CDEFG":
            self.assertGreaterEqual(ws.column_dimensions[letter].width, 22, letter)
        self.assertGreaterEqual(ws.column_dimensions["F"].width, 50)
        self.assertIsNone(ws.row_dimensions[2].height)
        wb.close()

    def test_full_sheet_column_definition_preserves_original_and_distant_widths(self):
        self.make_workbook()
        parts = self.package(self.source)
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cols = ET.Element("{" + NS["s"] + "}cols")
        ET.SubElement(cols, "{" + NS["s"] + "}col",
                      {"min": "1", "max": "16384", "width": "8.5", "customWidth": "1",
                       "hidden": "1", "style": "0", "outlineLevel": "1"})
        root.insert(list(root).index(root.find("s:sheetData", NS)), cols)
        self.replace_parts({"xl/worksheets/sheet1.xml": ET.tostring(root)})
        before = self.package(self.source)["xl/worksheets/sheet1.xml"]
        self.prepare()
        self.write_results()
        summary = self.export()
        after = self.package(summary["output"])["xl/worksheets/sheet1.xml"]
        for index in (1, 2, 8, 500, 16384):
            self.assertEqual(self.effective_column(before, index), self.effective_column(after, index), index)
        for index in range(3, 8):
            self.assertGreaterEqual(float(self.effective_column(after, index)["width"]), 22)
            self.assertNotEqual(self.effective_column(after, index).get("hidden"), "1")
        self.assertGreaterEqual(float(self.effective_column(after, 6)["width"]), 50)

    def test_multi_order_zero_amounts_and_all_fields_remain_aligned(self):
        self.make_workbook([["单号"], ["AB12345678;CD12345678"], ["AB12345678"]])
        self.prepare()
        self.write_results({
            "AB12345678": {"no": "AB12345678", "matched": 2, "data": [
                ["AB12345678", "2026-09-01", "0", "First & <complete> remark"],
                ["AB12345678", "2026-09-02", "12.30", "Second full remark"]]},
            "CD12345678": {"no": "CD12345678", "matched": 0,
                           "data": [["CD12345678", "", "", ""]]},
        })
        summary = self.export()
        wb = openpyxl.load_workbook(summary["output"], data_only=False)
        ws = wb.active
        self.assertEqual(ws["B2"].value, "AB12345678\nAB12345678\nCD12345678")
        self.assertEqual(ws["C2"].value, "2026-09-01\n2026-09-02\n")
        self.assertEqual(ws["D2"].value, "0\n12.30\n")
        self.assertEqual(ws["E2"].value, "First & <complete> remark\nSecond full remark\n")
        self.assertEqual(ws["F2"].value, "AB12345678: 匹配2条订单\nCD12345678: 未找到订单")
        self.assertEqual(ws["D3"].value, "0\n12.30")
        self.assertEqual(summary["multipleOrders"], 1)
        self.assertEqual(summary["notFound"], 1)
        wb.close()

    def test_source_and_existing_output_are_never_overwritten(self):
        self.make_workbook()
        self.prepare()
        self.write_results()
        source_hash = sha256(self.source)
        output = self.root / "existing.xlsx"
        output.write_bytes(b"existing output must survive")
        for protected in (self.source, output):
            with self.subTest(protected=protected.name), self.assertRaises(ValueError):
                self.export(protected)
        self.assertEqual(sha256(self.source), source_hash)
        self.assertEqual(output.read_bytes(), b"existing output must survive")
        self.assertFalse((self.work / "summary.json").exists())

    def test_source_change_after_prepare_is_rejected(self):
        self.make_workbook()
        self.prepare()
        self.write_results()
        self.replace_parts({"synthetic-added-after-prepare.txt": b"changed"})
        with self.assertRaises(ValueError):
            self.export()
        self.assertFalse((self.work / "summary.json").exists())
        self.assertFalse(list(self.work.glob("*.xlsx")))

    def test_incomplete_result_set_cannot_be_exported_as_missing_orders(self):
        self.make_workbook([["单号"], ["AB12345678"], ["CD12345678"]])
        self.prepare()
        self.write_results({"AB12345678": {"no": "AB12345678", "matched": 0,
                                            "data": [["AB12345678", "", "", ""]]}})
        with self.assertRaises(ValueError):
            self.export()
        self.assertFalse(list(self.work.glob("*.xlsx")))

    def test_prepare_refuses_existing_checkpoint(self):
        self.make_workbook()
        self.prepare()
        original = (self.work / "manifest.json").read_bytes()
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertEqual((self.work / "manifest.json").read_bytes(), original)

    def test_invalid_or_overlong_result_text_does_not_create_output(self):
        self.make_workbook()
        self.prepare()
        for remark in ("bad\x00text", "x" * 32768, "\U0001f600" * 16384):
            with self.subTest(utf16_units=len(remark.encode("utf-16-le")) // 2):
                self.write_results({"AB12345678": {"no": "AB12345678", "matched": 1,
                                                   "data": [["AB12345678", "", "0", remark]]}})
                with self.assertRaises(ValueError):
                    self.export()
                self.assertFalse(list(self.work.glob("*.xlsx")))
                self.assertFalse((self.work / "summary.json").exists())


if __name__ == "__main__":
    unittest.main()
