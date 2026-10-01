#!/usr/bin/env python3
"""Prepare tracking-number checkpoints and losslessly append order-query results.

Requires Python 3.10+ and openpyxl. No network/browser operations.
prepare --input FILE.xlsx --workdir DIR [--column HEADER] [--sheets A,B]
export --workdir DIR [--output FILE.xlsx]

The default column preference is 月结赔付单号, with exact-header fallback to
快递单号/物流单号/运单号/发货单号, then 单号 if none of these exist.
Explicit --column requires that exact header.
Headers are located in the first 50 actual rows; ambiguous matches fail closed.
Default traversal reports and skips sheets with no recognized header; explicitly
selected sheets must have a matching header. At least one tracking row is required.
results.json must map every manifest unique number to
{no: str, matched: int, data: [[no, confirm_date, paid_amount, remark], ...]}.
For unmatched numbers use matched=0 and exactly [[no, "", "", ""]].
Existing checkpoint/output files are never overwritten. Duplicate occurrences
are retained in their original rows; only the query list is deduplicated.
Prepared keys in unique and sheets[].rows[].nos are uppercased with one leading
@ removed. Extra boundary separators are ignored; original cells remain unchanged.
The runner may submit both
plain and @-prefixed search forms. Manifest metadata includes source, sha256,
unique, sheets, and original Excel row numbers. Export defaults to
WORKDIR/<source stem>_高阶版订单查询结果.xlsx; existing files are refused.
"""
import argparse
import copy
import hashlib
import json
import math
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from xml.dom import minidom as D

import openpyxl

HEADERS = ('月结赔付单号', '快递单号', '物流单号', '运单号', '发货单号', '单号')
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RID = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
RESULT_HEADERS = ['查询快递单号', '发货时间(confirm_date)', '实付金额(paid_amount)', '备注(remark)', '查询结果']
RESULT_WIDTHS = (24, 32, 30, 80, 32)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    h = hashlib.sha256()
    with open(path, 'rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save_new(path, value):
    with Path(path).open('x', encoding='utf-8', newline='\n') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write('\n')


def children(node, name):
    return [c for c in node.childNodes if c.nodeType == c.ELEMENT_NODE and c.localName == name]


def descendants(node, name):
    return list(node.getElementsByTagNameNS('*', name))


def resolve_part(base, target):
    path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(base), target))
    require(not path.startswith('../'), 'Invalid package relationship target')
    return path


def package_paths(z):
    roots = D.parseString(z.read('_rels/.rels'))
    offices = [r for r in descendants(roots, 'Relationship') if r.getAttribute('Type').endswith('/officeDocument') and r.getAttribute('TargetMode') != 'External']
    require(len(offices) == 1, 'Workbook relationship missing/ambiguous')
    wbpath = resolve_part('', offices[0].getAttribute('Target'))
    relpath = posixpath.join(posixpath.dirname(wbpath), '_rels', posixpath.basename(wbpath) + '.rels')
    rels = D.parseString(z.read(relpath))
    targets = {r.getAttribute('Id'): resolve_part(wbpath, r.getAttribute('Target')) for r in descendants(rels, 'Relationship') if r.getAttribute('TargetMode') != 'External'}
    workbook = D.parseString(z.read(wbpath))
    sheets = {s.getAttribute('name'): targets[s.getAttributeNS(RID, 'id')] for s in descendants(workbook, 'sheet')}
    return sheets


def colnum(ref):
    match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', ref)
    require(match is not None, 'Invalid cell reference: ' + ref)
    number = 0
    for c in match[1]:
        number = number * 26 + ord(c) - 64
    return number


def letter(number):
    result = ''
    while number:
        number, rem = divmod(number - 1, 26)
        result = chr(65 + rem) + result
    return result


def tracking_tokens(cell):
    value = cell.value
    if cell.data_type in ('f', 'e') or isinstance(value, bool):
        raise ValueError('Formula, error, or boolean is not a tracking number')
    if isinstance(value, (float, int)):
        require(not isinstance(value, float) or (math.isfinite(value) and value.is_integer()), 'Fractional/nonfinite numeric value')
        require(value >= 0 and len(str(int(value))) <= 15, 'Numeric tracking number may have lost Excel precision')
        text = str(int(value))
    else:
        require(isinstance(value, str), 'Unsupported tracking value type')
        text = value.strip()
    # Strip only known courier labels at token boundaries; leave unknown text for review.
    courier = r'(?:极兔(?:速递|快递)?|圆通(?:速递|快递)?|中通(?:快递|速递)?|申通(?:快递)?|韵达(?:快递)?|顺丰(?:速运|快递)?|京东(?:物流|快递)?|邮政(?:快递|速递)?|德邦(?:快递|物流)?|百世(?:快递)?)'
    text = re.sub(r'(^|[\s,，、;；|])' + courier + r'\s*[:：]?\s*(?=@?[A-Za-z0-9]{8,})', r'\1', text)
    # Empty boundary tokens carry no tracking information. Do not strip unknown
    # punctuation/annotations, repair numbers, or change the original cell.
    text = re.sub(r'^[\s,，、;；|]+|[\s,，、;；|]+$', '', text)
    parts = re.split(r'[\s,，、;；|]+', text)
    require(all(re.fullmatch(r'@?[A-Za-z0-9]{8,}', p) and re.search(r'[0-9]', p) for p in parts), 'Unrecognized tracking value; expected complete ASCII tracking numbers separated by spaces, commas, semicolons, or pipes')
    # Canonical keys also avoid PowerShell's case-insensitive hashtable collisions.
    # Preserve every occurrence in rows.nos; source cells themselves are untouched.
    return [p.removeprefix('@').upper() for p in parts]


def prepare(args):
    source = Path(args.input).resolve(strict=True)
    work = Path(args.workdir).resolve()
    work.mkdir(parents=True, exist_ok=True)
    require(source.suffix.lower() == '.xlsx', 'Input must be an .xlsx workbook')
    for name in ('manifest.json', 'results.json', 'prepare-errors.json'):
        require(not (work / name).exists(), 'Use a new workdir; checkpoint already exists: ' + name)
    original_hash = digest(source)
    sheets, unique, errors, skipped = [], {}, [], []
    with zipfile.ZipFile(source) as z:
        paths = package_paths(z)
        wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
        try:
            chosen = [s.strip() for s in args.sheets.split(',')] if args.sheets else wb.sheetnames
            require(len(chosen) == len(set(chosen)) and all(chosen), 'Duplicate or empty sheet selection')
            require(all(s in wb.sheetnames for s in chosen), 'Unknown sheet selection')
            for name in chosen:
                ws = wb[name]
                ws.reset_dimensions()  # producer metadata can incorrectly claim D238, etc.
                doc = D.parseString(z.read(paths[name]))
                actual_cells = descendants(doc, 'c')
                actual_rows = descendants(doc, 'row')
                maxrow = max([int(r.getAttribute('r')) for r in actual_rows] + [0])
                maxcol = max([colnum(c.getAttribute('r')) for c in actual_cells] + [0])
                rows = list(ws.iter_rows())
                candidates = []
                allowed = (args.column,) if args.column else HEADERS
                for rowid, row in enumerate(rows[:50], 1):
                    for column, cell in enumerate(row, 1):
                        if cell.value in allowed:
                            candidates.append((rowid, column, cell.value))
                preferred = [c for c in candidates if c[2] == HEADERS[0]] if args.column is None else []
                if preferred:
                    candidates = preferred
                elif args.column is None:
                    explicit = [c for c in candidates if c[2] != '单号']
                    if explicit:
                        candidates = explicit
                if not candidates and not args.sheets:
                    skipped.append({'name': name, 'reason': 'No recognized exact tracking header in first 50 rows'})
                    continue
                if len(candidates) != 1:
                    errors.append({'sheet': name, 'reason': 'Expected one exact tracking header', 'candidates': candidates})
                    continue
                headerrow, column, header = candidates[0]
                mapped = []
                for rowid, row in enumerate(rows, 1):
                    if rowid <= headerrow or len(row) < column:
                        continue
                    cell = row[column - 1]
                    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                        continue
                    try:
                        tokens = tracking_tokens(cell)
                    except ValueError as exc:
                        errors.append({'sheet': name, 'row': rowid, 'column': column, 'value': str(cell.value), 'reason': str(exc)})
                        continue
                    mapped.append({'row': rowid, 'nos': tokens})
                    for token in tokens:
                        unique[token] = None
                sheets.append({'name': name, 'headerRow': headerrow, 'column': column, 'header': header, 'maxRow': maxrow, 'maxCol': maxcol, 'rows': mapped})
        finally:
            wb.close()
    require(digest(source) == original_hash, 'Source changed during prepare')
    if not unique:
        errors.append({'reason': 'No tracking numbers found in selected workbook sheets'})
    if errors:
        save_new(work / 'prepare-errors.json', {'source': str(source), 'sha256': original_hash, 'errors': errors, 'skippedSheets': skipped})
        raise ValueError(f'Preparation failed with {len(errors)} issue(s); see {work / "prepare-errors.json"}. No manifest was created.')
    require(bool(unique), 'No tracking numbers found')
    manifest = {'schemaVersion': 1, 'source': str(source), 'sha256': original_hash, 'sheets': sheets, 'unique': list(unique), 'skippedSheets': skipped}
    save_new(work / 'manifest.json', manifest)
    return {'manifest': str(work / 'manifest.json'), 'unique': len(unique), 'skippedSheets': skipped, 'sheets': [{'name': s['name'], 'rows': len(s['rows']), 'column': s['column'], 'header': s['header'], 'maxRow': s['maxRow'], 'maxCol': s['maxCol']} for s in sheets]}


def make_element(doc, root, name, attrs=None):
    node = doc.createElementNS(root.namespaceURI, (root.prefix + ':' if root.prefix else '') + name)
    for key, value in (attrs or {}).items():
        node.setAttribute(key, str(value))
    return node


def column_settings(root):
    """Read effective explicit formatting, rejecting ambiguous column ranges."""
    blocks = children(root, 'cols')
    require(len(blocks) <= 1, 'Multiple column definition blocks unsupported')
    settings = {}
    for block in blocks:
        for col in children(block, 'col'):
            first, last = int(col.getAttribute('min')), int(col.getAttribute('max'))
            require(1 <= first <= last <= 16384, 'Invalid column definition range')
            attrs = {key: value for key, value in col.attributes.items() if key not in ('min', 'max')}
            for number in range(first, last + 1):
                require(number not in settings, 'Overlapping column definitions unsupported')
                settings[number] = attrs
    return settings


def set_result_widths(doc, root, start):
    """Set widths only for appended columns; split spanning defaults if needed."""
    column_settings(root)
    blocks = children(root, 'cols')
    if blocks:
        cols = blocks[0]
    else:
        cols = make_element(doc, root, 'cols')
        root.insertBefore(cols, children(root, 'sheetData')[0])
    end = start + len(RESULT_WIDTHS) - 1
    for col in children(cols, 'col'):
        first, last = int(col.getAttribute('min')), int(col.getAttribute('max'))
        if first > end or last < start:
            continue
        # A producer may apply a single width/style to A:XFD. Keep its exact
        # attributes everywhere except the five newly populated columns.
        for left, right in ((first, start - 1), (end + 1, last)):
            if left <= right:
                part = col.cloneNode(deep=True)
                part.setAttribute('min', str(left))
                part.setAttribute('max', str(right))
                cols.insertBefore(part, col)
        cols.removeChild(col)
    for offset, width in enumerate(RESULT_WIDTHS):
        number = start + offset
        col = make_element(doc, root, 'col', {
            'min': number, 'max': number, 'width': width, 'customWidth': 1,
        })
        following = next((c for c in children(cols, 'col') if int(c.getAttribute('min')) > number), None)
        cols.insertBefore(col, following)


def verify_and_restore_columns(olddoc, newdoc, start):
    """Verify widths and unchanged original formatting before structural diff."""
    oldroot, newroot = olddoc.documentElement, newdoc.documentElement
    old, new = column_settings(oldroot), column_settings(newroot)
    result_columns = set(range(start, start + len(RESULT_WIDTHS)))
    require(all(old.get(c) == new.get(c) for c in (old.keys() | new.keys()) - result_columns),
            'Original column formatting changed')
    for offset, width in enumerate(RESULT_WIDTHS):
        require(new.get(start + offset) == {'width': str(width), 'customWidth': '1'},
                'Result column width verification failed')
    oldblocks, newblocks = children(oldroot, 'cols'), children(newroot, 'cols')
    if oldblocks:
        # Preserve comments, extension content and block attributes, too.
        oldblock, newblock = oldblocks[0].cloneNode(deep=True), newblocks[0].cloneNode(deep=True)
        for block in (oldblock, newblock):
            for col in children(block, 'col'):
                block.removeChild(col)
        require(oldblock.toxml() == newblock.toxml(), 'Column definition metadata changed')
        newroot.replaceChild(newdoc.importNode(oldblocks[0], True), newblocks[0])
    else:
        newroot.removeChild(newblocks[0])


def validate_results(manifest, results):
    require(isinstance(results, dict) and set(results) == set(manifest['unique']), 'Results do not exactly cover manifest unique numbers')
    for no, item in results.items():
        require(isinstance(item, dict) and item.get('no') == no, 'Result number mismatch: ' + no)
        count, records = item.get('matched'), item.get('data')
        require(type(count) is int and count >= 0 and isinstance(records, list), 'Invalid result count/data: ' + no)
        require(len(records) == max(1, count), 'Result row count mismatch: ' + no)
        for record in records:
            require(isinstance(record, list) and len(record) == 4 and record[0] == no and all(isinstance(v, str) for v in record), 'Malformed result data: ' + no)
        require(count != 0 or records == [[no, '', '', '']], 'Unmatched result must contain blank fields: ' + no)
    expected_unique = list(dict.fromkeys(no for sheet in manifest['sheets'] for row in sheet['rows'] for no in row['nos']))
    require(expected_unique == manifest['unique'], 'Manifest row mapping does not match unique checkpoint')


def export(args):
    work = Path(args.workdir).resolve()
    manifest, results = load(work / 'manifest.json'), load(work / 'results.json')
    source = Path(manifest['source']).resolve(strict=True)
    output = Path(args.output).resolve() if args.output else work / (source.stem + '_高阶版订单查询结果.xlsx')
    require(output != source, 'Output must never be the source workbook')
    require(output.suffix.lower() == '.xlsx', 'Output must be .xlsx')
    require(not output.exists(), 'Refusing to overwrite output: ' + str(output))
    require(not (work / 'summary.json').exists(), 'Refusing to overwrite summary.json')
    require(digest(source) == manifest['sha256'], 'Source hash differs from prepare checkpoint')
    validate_results(manifest, results)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = None
    try:
        with zipfile.ZipFile(source) as zin:
            require(len(zin.namelist()) == len(set(zin.namelist())), 'Duplicate ZIP entries unsupported')
            paths = package_paths(zin)
            changes, expected, reports, result_starts = {}, {}, [], {}
            for sheet in manifest['sheets']:
                path = paths[sheet['name']]
                doc = D.parseString(zin.read(path))
                root = doc.documentElement
                data = children(root, 'sheetData')[0]
                rows = {int(r.getAttribute('r')): r for r in children(data, 'row')}
                existing = {c.getAttribute('r'): c for r in rows.values() for c in children(r, 'c')}
                # Append beyond ALL existing cell nodes, including blank styled cells.
                # Original cells/styles stay intact; widths change only in new columns.
                lastcol = max([colnum(ref) for ref in existing] + [0])
                for merge in descendants(root, 'mergeCell'):
                    lastcol = max(lastcol, colnum(merge.getAttribute('ref').split(':')[-1]))
                start = lastcol + 1
                require(start + 4 <= 16384, 'No space for result columns')
                result_starts[sheet['name']] = start
                set_result_widths(doc, root, start)
                values = {sheet['headerRow']: RESULT_HEADERS}
                counts = Counter()
                seen_rows = set()
                for row in sheet['rows']:
                    rowid = row['row']
                    require(rowid not in seen_rows and rowid in rows and rowid > sheet['headerRow'], 'Invalid/duplicate manifest row')
                    seen_rows.add(rowid)
                    records = [record for no in row['nos'] for record in results[no]['data']]
                    statuses = []
                    for no in row['nos']:
                        count = results[no]['matched']
                        state = '未找到订单' if count == 0 else ('匹配成功' if count == 1 else f'匹配{count}条订单')
                        statuses.append(f'{no}: {state}' if len(row['nos']) > 1 else state)
                    values[rowid] = ['\n'.join(r[i] for r in records) for i in range(4)] + ['\n'.join(statuses)]
                    counts['rows'] += 1
                    counts['allMatched' if all(results[n]['matched'] for n in row['nos']) else 'hasNotFound'] += 1
                    if any(results[n]['matched'] > 1 for n in row['nos']):
                        counts['multipleOrders'] += 1
                expected[sheet['name']] = {}
                for rowid, fields in values.items():
                    require(rowid in rows, 'Header row missing from original XML')
                    for offset, value in enumerate(fields):
                        require(len(value.encode('utf-16-le')) // 2 <= 32767, 'Result exceeds Excel cell text limit')
                        require(not re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value), 'Result contains invalid XML control characters')
                        coord = f'{letter(start + offset)}{rowid}'
                        require(coord not in existing, 'Attempt to overwrite original cell: ' + coord)
                        cell = make_element(doc, root, 'c', {'r': coord, 't': 'inlineStr'})
                        inline = make_element(doc, root, 'is')
                        text = make_element(doc, root, 't', {'xml:space': 'preserve'})
                        text.appendChild(doc.createTextNode(value))
                        inline.appendChild(text)
                        cell.appendChild(inline)
                        # c elements precede possible row extension elements.
                        trailing = next((n for n in rows[rowid].childNodes if n.nodeType == n.ELEMENT_NODE and n.localName != 'c'), None)
                        rows[rowid].insertBefore(cell, trailing)
                        expected[sheet['name']][coord] = value
                dims = children(root, 'dimension')
                if dims:
                    dim = dims[0]
                else:
                    dim = make_element(doc, root, 'dimension')
                    before = next((n for n in root.childNodes if n.nodeType == n.ELEMENT_NODE and n.localName != 'sheetPr'), None)
                    root.insertBefore(dim, before)
                dim.setAttribute('ref', f'A1:{letter(start + 4)}{max(rows)}')
                changes[path] = doc.toxml(encoding='utf-8')
                reports.append({'sheet': sheet['name'], 'resultColumns': f'{letter(start)}:{letter(start + 4)}', 'resultColumnWidths': list(RESULT_WIDTHS), **counts})
            fd, tempname = tempfile.mkstemp(prefix='excel-export-', suffix='.xlsx', dir=output.parent)
            os.close(fd)
            temp = Path(tempname)
            with zipfile.ZipFile(temp, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
                zout.comment = zin.comment
                for info in zin.infolist():
                    zout.writestr(copy.copy(info), changes.get(info.filename, zin.read(info.filename)))
            with zipfile.ZipFile(temp) as zout:
                require(zin.namelist() == zout.namelist(), 'ZIP part names changed')
                for name in zin.namelist():
                    if name not in changes:
                        require(zin.read(name) == zout.read(name), 'Unrelated package part changed: ' + name)
                for sheet in manifest['sheets']:
                    path = paths[sheet['name']]
                    olddoc, newdoc = D.parseString(zin.read(path)), D.parseString(zout.read(path))
                    newcells = {c.getAttribute('r'): c for c in descendants(newdoc, 'c')}
                    for old in descendants(olddoc, 'c'):
                        require(old.toxml() == newcells[old.getAttribute('r')].toxml(), 'Original cell changed: ' + old.getAttribute('r'))
                    # Check effective formatting outside the added result columns,
                    # then restore column definitions for the exact structural diff.
                    verify_and_restore_columns(olddoc, newdoc, result_starts[sheet['name']])
                    # Remove intentional additions and restore dimension; the rest
                    # of the original worksheet must then be identical.
                    for coord in expected[sheet['name']]:
                        node = newcells[coord]
                        node.parentNode.removeChild(node)
                    olddim, newdim = children(olddoc.documentElement, 'dimension'), children(newdoc.documentElement, 'dimension')
                    if olddim:
                        newdim[0].parentNode.replaceChild(newdoc.importNode(olddim[0], True), newdim[0])
                    else:
                        newdim[0].parentNode.removeChild(newdim[0])
                    require(olddoc.toxml() == newdoc.toxml(), 'Original worksheet structure changed')
            wb = openpyxl.load_workbook(temp, read_only=True, data_only=False)
            try:
                for name, cells in expected.items():
                    ws = wb[name]
                    ws.reset_dimensions()
                    actual = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
                    require(all((actual.get(coord) or '') == value for coord, value in cells.items()), 'Exported result cell verification failed: ' + name)
            finally:
                wb.close()
        require(digest(source) == manifest['sha256'], 'Source changed during export')
        # Exclusive creation avoids races and never truncates an existing output.
        with output.open('xb') as target, temp.open('rb') as original:
            for block in iter(lambda: original.read(1024 * 1024), b''):
                target.write(block)
        summary = {'output': str(output), 'source': str(source), 'sourceSha256': manifest['sha256'], 'unique': len(results), 'matched': sum(v['matched'] > 0 for v in results.values()), 'notFound': sum(v['matched'] == 0 for v in results.values()), 'multipleOrders': sum(v['matched'] > 1 for v in results.values()), 'sheets': reports, 'skippedSheets': manifest.get('skippedSheets', []), 'verifiedOriginalCells': True, 'verifiedOriginalWorksheetStructure': True, 'verifiedUnchangedPackageParts': True, 'verifiedResultCells': True, 'sourceUnchanged': True}
        save_new(work / 'summary.json', summary)
        return summary
    finally:
        if temp is not None:
            temp.unlink(missing_ok=True)


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest='command', required=True)
    prep = commands.add_parser('prepare')
    prep.add_argument('--input', required=True)
    prep.add_argument('--workdir', required=True)
    prep.add_argument('--column', help='Exact header; default preference 月结赔付单号 with approved header fallback')
    prep.add_argument('--sheets', help='Optional comma-separated sheet names; default all sheets')
    exp = commands.add_parser('export')
    exp.add_argument('--workdir', required=True)
    exp.add_argument('--output')
    args = parser.parse_args()
    try:
        print(json.dumps(prepare(args) if args.command == 'prepare' else export(args), ensure_ascii=False))
    except (ValueError, OSError, KeyError, zipfile.BadZipFile) as exc:
        print(json.dumps({'ok': False, 'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
